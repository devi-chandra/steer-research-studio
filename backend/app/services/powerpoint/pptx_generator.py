import os
import json
from pathlib import Path
from typing import Dict, List, Any
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from app.core.config import settings
from app.core.database import get_db_connection

# Steer Brand Palette RGB
COLOR_CREAM = RGBColor(245, 235, 223)    # #F5EBDF
COLOR_MAROON = RGBColor(102, 20, 20)     # #661414
COLOR_RUST = RGBColor(192, 78, 42)       # #C04E2A
COLOR_INK = RGBColor(53, 29, 29)         # #351D1D
COLOR_CLAY = RGBColor(181, 135, 119)     # #B58777
COLOR_BLUSH = RGBColor(230, 203, 189)    # #E6CBBD
COLOR_WHITE = RGBColor(255, 255, 255)

def _apply_slide_background(slide, prs):
    # Add full slide background shape with Cream color
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_CREAM
    bg.line.color.rgb = COLOR_CREAM

    # Top brand bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.18))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_MAROON
    bar.line.color.rgb = COLOR_MAROON

def _add_slide_header(slide, title_text: str, badge_text: str = "STEER MEDICAL STUDIO"):
    # Header badge
    tb_b = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(8), Inches(0.3))
    p_b = tb_b.text_frame.paragraphs[0]
    p_b.text = badge_text.upper()
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_RUST

    # Title
    tb_t = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.5), Inches(0.8))
    p_t = tb_t.text_frame.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_MAROON

    # Rust underline
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(2.5), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_RUST
    line.line.color.rgb = COLOR_RUST

def _add_provenance_footer(slide, claim_id: str, page_num: int, doi: str):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.5), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = f"Source provenance: Claim {claim_id} • Source Page {page_num} • DOI: {doi or 'Not available'} • Human review status: pending"
    p.font.size = Pt(9)
    p.font.color.rgb = COLOR_CLAY

def generate_powerpoint_presentation(project_id: str) -> Dict[str, Any]:
    """
    Feature 4: Research-to-PowerPoint Studio.
    Generates genuine, editable 10-slide PowerPoint (.pptx) deck in Steer clinical editorial theme.
    """
    conn = get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT title, doi, year, authors FROM projects WHERE id = ?", (project_id,))
    proj = cursor.fetchone()
    if not proj:
        conn.close()
        raise ValueError("Project not found.")

    cursor.execute("""
    SELECT claim_id, claim_text, exact_quote, page_number, section_heading,
           numerical_result, sample_size, limitations, uncertainty, evidence_status
    FROM claims WHERE project_id = ? ORDER BY page_number ASC
    """, (project_id,))
    claims = [dict(r) for r in cursor.fetchall()]
    conn.close()

    paper_title = proj["title"]
    paper_doi = proj["doi"] or "Not available"
    paper_year = proj["year"] or "Not available"
    paper_authors = proj["authors"] or "Not available"

    prs = Presentation()
    prs.slide_width = Inches(13.33)  # 16:9 Widescreen standard
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # --- SLIDE 1: Title & Provenance ---
    s1 = prs.slides.add_slide(blank_layout)
    _apply_slide_background(s1, prs)

    # Accent decorative box
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.5), Inches(11.33), Inches(4.8))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_WHITE
    card.line.color.rgb = COLOR_BLUSH
    card.line.width = Pt(1.5)

    tb1 = s1.shapes.add_textbox(Inches(1.5), Inches(1.9), Inches(10.33), Inches(1.8))
    p = tb1.text_frame.paragraphs[0]
    p.text = "CLINICAL RESEARCH BRIEFING"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_RUST

    p2 = tb1.text_frame.add_paragraph()
    p2.text = paper_title
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_MAROON

    tb_meta = s1.shapes.add_textbox(Inches(1.5), Inches(3.9), Inches(10.33), Inches(1.8))
    pm1 = tb_meta.text_frame.paragraphs[0]
    pm1.text = f"Investigators: {paper_authors}"
    pm1.font.size = Pt(14)
    pm1.font.color.rgb = COLOR_INK

    pm2 = tb_meta.text_frame.add_paragraph()
    pm2.text = f"Publication Year: {paper_year} • DOI: {paper_doi}"
    pm2.font.size = Pt(14)
    pm2.font.color.rgb = COLOR_INK

    pm3 = tb_meta.text_frame.add_paragraph()
    pm3.text = "Evidence extraction record • Human medical review status: pending"
    pm3.font.size = Pt(12)
    pm3.font.color.rgb = COLOR_CLAY

    # --- SLIDE 2: Background & Clinical Question ---
    s2 = prs.slides.add_slide(blank_layout)
    _apply_slide_background(s2, prs)
    _add_slide_header(s2, "Background & Research Objective", "CLINICAL CONTEXT")

    card2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.8))
    card2.fill.solid()
    card2.fill.fore_color.rgb = COLOR_WHITE
    card2.line.color.rgb = COLOR_BLUSH

    tb2 = s2.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "Primary Research Objective:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    p_body = tf2.add_paragraph()
    p_body.text = f"Source title: {paper_title[:300]}"
    p_body.font.size = Pt(15)
    p_body.font.color.rgb = COLOR_INK

    p_why = tf2.add_paragraph()
    p_why.text = "\nClinical Significance:"
    p_why.font.size = Pt(16)
    p_why.font.bold = True
    p_why.font.color.rgb = COLOR_RUST

    p_whyb = tf2.add_paragraph()
    c0 = claims[0] if claims else {}
    p_whyb.text = f"Source excerpt: '{c0.get('exact_quote') or 'No extracted evidence is available for this section.'}'"
    p_whyb.font.size = Pt(14)
    p_whyb.font.color.rgb = COLOR_INK

    _add_provenance_footer(s2, c0.get("claim_id", "CLM-01"), c0.get("page_number", 1), paper_doi)

    # --- SLIDE 3: Study Design & Methodology ---
    s3 = prs.slides.add_slide(blank_layout)
    _apply_slide_background(s3, prs)
    _add_slide_header(s3, "Study Design & Investigative Methodology", "METHODOLOGY")

    card3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.8))
    card3.fill.solid()
    card3.fill.fore_color.rgb = COLOR_WHITE
    card3.line.color.rgb = COLOR_BLUSH

    tb3 = s3.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.0))
    tf3 = tb3.text_frame
    tf3.word_wrap = True

    c1 = claims[1] if len(claims) > 1 else (claims[0] if claims else {})
    p = tf3.paragraphs[0]
    p.text = "Trial Architecture:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    p_body = tf3.add_paragraph()
    p_body.text = f"Source excerpt: {c1.get('exact_quote') or 'No extracted methods evidence is available.'}\n\nStudy classification, protocol environment, and endpoints are not shown unless directly extracted from the source."
    p_body.font.size = Pt(15)
    p_body.font.color.rgb = COLOR_INK

    _add_provenance_footer(s3, c1.get("claim_id", "CLM-01"), c1.get("page_number", 1), paper_doi)

    # --- SLIDE 4: Cohort & Participants ---
    s4 = prs.slides.add_slide(blank_layout)
    _apply_slide_background(s4, prs)
    _add_slide_header(s4, "Participants & Cohort Characteristics", "COHORT DEMOGRAPHICS")

    card4 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.8))
    card4.fill.solid()
    card4.fill.fore_color.rgb = COLOR_WHITE
    card4.line.color.rgb = COLOR_BLUSH

    tb4 = s4.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.0))
    tf4 = tb4.text_frame
    tf4.word_wrap = True

    p = tf4.paragraphs[0]
    p.text = f"Sample Size: {c0.get('sample_size') or 'Not extracted'}"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = COLOR_RUST

    p_pop = tf4.add_paragraph()
    p_pop.text = f"\nPopulation:\n{c0.get('population') or 'Not extracted'}"
    p_pop.font.size = Pt(15)
    p_pop.font.color.rgb = COLOR_INK

    p_note = tf4.add_paragraph()
    p_note.text = "\nDemographic details are included only when present in extracted source evidence."
    p_note.font.size = Pt(13)
    p_note.font.color.rgb = COLOR_CLAY

    _add_provenance_footer(s4, c0.get("claim_id", "CLM-01"), c0.get("page_number", 1), paper_doi)

    # --- SLIDE 5: Key Findings ---
    s5 = prs.slides.add_slide(blank_layout)
    _apply_slide_background(s5, prs)
    _add_slide_header(s5, "Primary Findings & Evidence Synthesis", "KEY FINDINGS")

    card5 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.8))
    card5.fill.solid()
    card5.fill.fore_color.rgb = COLOR_WHITE
    card5.line.color.rgb = COLOR_BLUSH

    tb5 = s5.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.0))
    tf5 = tb5.text_frame
    tf5.word_wrap = True

    c2 = claims[2] if len(claims) > 2 else (claims[0] if claims else {})
    p = tf5.paragraphs[0]
    p.text = f"Finding: {c2.get('claim_text') or 'No finding was extracted from the source.'}"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    p_q = tf5.add_paragraph()
    p_q.text = f"\nVerbatim Supporting Quotation:\n\"{c2.get('exact_quote', '')}\""
    p_q.font.size = Pt(14)
    p_q.font.color.rgb = COLOR_INK

    p_s = tf5.add_paragraph()
    p_s.text = f"\nEvidence Status: {c2.get('evidence_status') or 'Not available'} • Human medical review: pending"
    p_s.font.size = Pt(13)
    p_s.font.bold = True
    p_s.font.color.rgb = COLOR_RUST

    _add_provenance_footer(s5, c2.get("claim_id", "CLM-01"), c2.get("page_number", 1), paper_doi)

    # --- SLIDE 6: Structured Evidence Table ---
    s6 = prs.slides.add_slide(blank_layout)
    _apply_slide_background(s6, prs)
    _add_slide_header(s6, "Quantitative Evidence Ledger", "MEASUREMENTS & METRICS")

    # Add real PowerPoint Table
    rows = min(5, len(claims) + 1)
    cols = 4
    left = Inches(0.8)
    top = Inches(1.8)
    width = Inches(11.7)
    height = Inches(4.5)

    table_shape = s6.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.columns[0].width = Inches(1.5)
    table.columns[1].width = Inches(4.5)
    table.columns[2].width = Inches(3.2)
    table.columns[3].width = Inches(2.5)

    headers = ["Claim ID", "Statement & Outcome", "Numerical Result", "Evidence Status"]
    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_MAROON
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

    for row_idx in range(1, rows):
        claim_item = claims[row_idx - 1]
        c0 = table.cell(row_idx, 0)
        c0.text = claim_item.get("claim_id", f"CLM-{row_idx:02d}")
        c1 = table.cell(row_idx, 1)
        c1.text = claim_item.get("claim_text", "")[:95]
        c2 = table.cell(row_idx, 2)
        c2.text = claim_item.get("numerical_result", "Reported in text")
        c3 = table.cell(row_idx, 3)
        c3.text = claim_item.get("evidence_status", "Directly supported")

        for c_cell in [c0, c1, c2, c3]:
            for p in c_cell.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.color.rgb = COLOR_INK

    _add_provenance_footer(s6, "CLM-ALL", 1, paper_doi)

    # --- SLIDE 7: Interpretation & Mechanism ---
    s7 = prs.slides.add_slide(blank_layout)
    _apply_slide_background(s7, prs)
    _add_slide_header(s7, "Interpretation & Mechanism", "SOURCE EVIDENCE")

    card7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.8))
    card7.fill.solid()
    card7.fill.fore_color.rgb = COLOR_WHITE
    card7.line.color.rgb = COLOR_BLUSH

    tb7 = s7.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.0))
    tf7 = tb7.text_frame
    tf7.word_wrap = True
    mechanism_claim = claims[2] if len(claims) > 2 else (claims[0] if claims else {})

    p = tf7.paragraphs[0]
    p.text = "Source-reported interpretation:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    p_m = tf7.add_paragraph()
    p_m.text = f"{mechanism_claim.get('exact_quote') or 'No mechanism or interpretation was extracted.'}\n\nNo biological mechanism is inferred when the source does not report one."
    p_m.font.size = Pt(14)
    p_m.font.color.rgb = COLOR_INK

    _add_provenance_footer(s7, "CLM-02", 1, paper_doi)

    # --- SLIDE 8: Limitations & Medical Caveats ---
    s8 = prs.slides.add_slide(blank_layout)
    _apply_slide_background(s8, prs)
    _add_slide_header(s8, "Limitations, Uncertainty & Clinical Bounds", "METHODOLOGICAL BOUNDS")

    card8 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.8))
    card8.fill.solid()
    card8.fill.fore_color.rgb = COLOR_WHITE
    card8.line.color.rgb = COLOR_MAROON
    card8.line.width = Pt(2.0)

    tb8 = s8.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.0))
    tf8 = tb8.text_frame
    tf8.word_wrap = True

    c3 = claims[3] if len(claims) > 3 else (claims[0] if claims else {})
    p = tf8.paragraphs[0]
    p.text = "CRITICAL STUDY LIMITATIONS:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    p_lim = tf8.add_paragraph()
    p_lim.text = f"• Directly stated limitations: {c3.get('limitations') or 'Not extracted'}\n• Statistical uncertainty: {c3.get('uncertainty') or 'Not extracted'}\n• This research summary is informational and does not provide personal medical advice."
    p_lim.font.size = Pt(14)
    p_lim.font.color.rgb = COLOR_INK

    _add_provenance_footer(s8, c3.get("claim_id", "CLM-01"), c3.get("page_number", 1), paper_doi)

    # --- SLIDE 9: Practical Clinical Takeaways ---
    s9 = prs.slides.add_slide(blank_layout)
    _apply_slide_background(s9, prs)
    _add_slide_header(s9, "Practical Takeaways & Future Directions", "CLINICAL TRANSLATION")

    card9 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.8))
    card9.fill.solid()
    card9.fill.fore_color.rgb = COLOR_WHITE
    card9.line.color.rgb = COLOR_BLUSH

    tb9 = s9.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.0))
    tf9 = tb9.text_frame
    tf9.word_wrap = True

    p = tf9.paragraphs[0]
    p.text = "Key Takeaways for Clinicians & Healthcare Teams:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    p_t = tf9.add_paragraph()
    p_t.text = f"Source-supported takeaway: {mechanism_claim.get('exact_quote') or 'No evidence-backed takeaway was extracted.'}\n\nThis deck does not make clinical recommendations."
    p_t.font.size = Pt(14)
    p_t.font.color.rgb = COLOR_INK

    _add_provenance_footer(s9, "CLM-REF", 1, paper_doi)

    # --- SLIDE 10: References & Steer Verification Ledger ---
    s10 = prs.slides.add_slide(blank_layout)
    _apply_slide_background(s10, prs)
    _add_slide_header(s10, "Source Provenance & Verification Audit", "REFERENCES")

    card10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.8))
    card10.fill.solid()
    card10.fill.fore_color.rgb = COLOR_WHITE
    card10.line.color.rgb = COLOR_BLUSH

    tb10 = s10.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.0))
    tf10 = tb10.text_frame
    tf10.word_wrap = True

    p = tf10.paragraphs[0]
    p.text = "Bibliographic & Verification Record:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_MAROON

    p_ref = tf10.add_paragraph()
    p_ref.text = f"Primary Source Document: {paper_title}\nDOI: {paper_doi}\nExtracted claims: {len(claims)} records\nHuman medical review status: pending"
    p_ref.font.size = Pt(14)
    p_ref.font.color.rgb = COLOR_INK

    _add_provenance_footer(s10, "CLM-AUDIT", 1, paper_doi)

    # Save PPTX file
    out_dir = settings.GENERATED_DIR / "pptx"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_file = out_dir / f"presentation_{project_id}.pptx"
    prs.save(str(out_file))

    # Verify structural validity by reopening
    test_prs = Presentation(str(out_file))
    verified_slide_count = len(test_prs.slides)

    # Save asset record in DB
    asset_id = f"ast-pptx-{project_id}"
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
    INSERT OR REPLACE INTO assets (id, project_id, asset_type, title, status, file_path, preview_url, metadata_json, updated_at)
    VALUES (?, ?, 'pptx', 'Research-to-PowerPoint Deck (10 Slides)', 'rendered', ?, ?, ?, CURRENT_TIMESTAMP)
    """, (
        asset_id,
        project_id,
        str(out_file),
        f"/storage/generated/pptx/presentation_{project_id}.pptx",
        json.dumps({
            "slide_count": verified_slide_count,
            "theme": "Steer Clinical Editorial",
            "file_size": out_file.stat().st_size
        })
    ))
    conn.commit()
    conn.close()

    return {
        "project_id": project_id,
        "asset_id": asset_id,
        "slide_count": verified_slide_count,
        "file_path": str(out_file),
        "download_url": f"/storage/generated/pptx/presentation_{project_id}.pptx",
        "file_size_bytes": out_file.stat().st_size
    }
